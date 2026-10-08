import json
from io import BytesIO

import pytest
from docx import Document

from app.core.llm_fallback import FallbackLLMManager
from app.graphs.cv_screening_graph import build_cv_screening_graph
from app.services.embeddings import HashingEmbedder
from tests.checklist_factory import sample_queries

CV_TEXT = (
    "Ada Lovelace built production FastAPI services and shipped a policy-compliant hiring platform "
    "with Docker. She is a female engineer."
)


class ScriptedProvider:
    def __init__(self) -> None:
        self.calls = 0

    async def complete(self, prompt: str, *, system: str | None = None) -> str:
        del system
        self.calls += 1
        if "Extract candidate attributes" in prompt:
            return json.dumps(
                {
                    "full_name": "Ada Lovelace",
                    "email": "ada@example.com",
                    "skills": ["FastAPI"],
                    "work_history": ["Sift"],
                    "education": ["Mathematics"],
                    "experience_years": 4,
                }
            )
        evidence = "shipped a policy-compliant hiring platform"
        query_id = "Q1"
        for index in range(1, 11):
            marker = f"Query: Q{index} "
            if marker in prompt:
                query_id = f"Q{index}"
                break
        return json.dumps(
            {"query_id": query_id, "status": "met", "confidence": 0.9, "evidence": evidence}
        )


def _docx() -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_paragraph(CV_TEXT)
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_workflow_scores_a_cv_against_the_locked_checklist() -> None:
    provider = ScriptedProvider()
    run = build_cv_screening_graph(FallbackLLMManager(provider, provider, provider), HashingEmbedder())
    result = await run(
        {
            "filename": "ada.docx",
            "file_bytes": _docx(),
            "job_description": "A" * 120,
            "locked_checklist": sample_queries(),
        }
    )
    assert result["parse_method"] == "native"
    assert result["aggregate_score"] == 1
    assert result["extracted_attributes"]["full_name"] == "Ada Lovelace"
    assert "female" in result["fairness_flags"]
    assert len(result["query_evaluations"]) == 10
    assert provider.calls == 11
    assert result["justifications"]
    assert result["extracted_attributes"]["experience_years"] == 4.0


@pytest.mark.asyncio
async def test_workflow_uses_computed_tenure_and_keeps_pgvector_evidence() -> None:
    provider = TenureProvider()
    run = build_cv_screening_graph(FallbackLLMManager(provider, provider, provider), HashingEmbedder())
    checklist = sample_queries()
    checklist[0] = {
        **checklist[0],
        "question": "Does the candidate have experience with pgvector?",
        "dimension": "technical",
        "evidence_hints": ["pgvector"],
    }
    checklist[6] = {
        **checklist[6],
        "question": "Does the candidate have at least 8 years of professional experience?",
        "dimension": "experience",
        "evidence_hints": ["experience"],
    }
    result = await run(
        {
            "filename": "kaji.docx",
            "file_bytes": _docx_with(TENURE_CV),
            "job_description": "A" * 120,
            "locked_checklist": checklist,
        }
    )
    by_query = {item["query_id"]: item for item in result["query_evaluations"]}
    assert result["extracted_attributes"]["experience_years"] >= 11
    assert by_query["Q1"]["status"] == "met"
    assert "pgvector" in str(by_query["Q1"]["evidence"]).casefold()
    assert by_query["Q7"]["status"] == "met"
    assert "2013" in str(by_query["Q7"]["evidence"])


TENURE_CV = (
    "Kaji Sadman Sakib shipped a policy-compliant hiring platform. "
    "Engineer at Northwind, 2013 to 2016. "
    "Engineer at Contoso, 2016 to 2019. "
    "Engineer at Fabrikam, 2019 to 2023, built pgvector search. "
    "Engineer at Litware, 2023 to 2026."
)


class TenureProvider:
    async def complete(self, prompt: str, *, system: str | None = None) -> str:
        del system
        if "Extract candidate attributes" in prompt:
            return json.dumps(
                {
                    "full_name": "Kaji Sadman Sakib",
                    "email": "kaji@example.com",
                    "skills": ["pgvector", "FastAPI"],
                    "work_history": [
                        {
                            "title": "Engineer",
                            "company": "Northwind",
                            "start": "2013-01",
                            "end": "2016-01",
                            "summary": "APIs",
                        },
                        {
                            "title": "Engineer",
                            "company": "Contoso",
                            "start": "2016-01",
                            "end": "2019-06",
                            "summary": "Search",
                        },
                        {
                            "title": "Engineer",
                            "company": "Fabrikam",
                            "start": "2019-06",
                            "end": "2023-01",
                            "summary": "pgvector search",
                        },
                        {
                            "title": "Engineer",
                            "company": "Litware",
                            "start": "2023-01",
                            "end": "2026-01",
                            "summary": "FastAPI",
                        },
                    ],
                    "education": ["B.Sc."],
                    "experience_years": 3,
                }
            )
        if "experience with pgvector" in prompt:
            return json.dumps(
                {
                    "query_id": "Q1",
                    "status": "insufficient_evidence",
                    "confidence": 0.2,
                    "evidence": "deep experience with PG Vector",
                }
            )
        if "professional experience" in prompt:
            return json.dumps(
                {
                    "query_id": "Q7",
                    "status": "unmet",
                    "confidence": 0.4,
                    "evidence": "only three years are visible",
                }
            )
        evidence = "shipped a policy-compliant hiring platform"
        query_id = "Q1"
        for index in range(1, 11):
            marker = f"Query: Q{index} "
            if marker in prompt:
                query_id = f"Q{index}"
                break
        return json.dumps({"query_id": query_id, "status": "met", "confidence": 0.9, "evidence": evidence})


def _docx_with(text: str) -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


PYTHON_CV = (
    "Ada Lovelace has expertise in FastAPI and shipped a policy-compliant hiring platform."
)


class CloudScreenProvider:
    def __init__(self) -> None:
        self.calls = 0

    async def complete(self, prompt: str, *, system: str | None = None) -> str:
        del system
        self.calls += 1
        if "Extract candidate attributes" in prompt:
            return json.dumps(
                {
                    "full_name": "Ada Lovelace",
                    "email": "ada@example.com",
                    "skills": ["FastAPI"],
                    "work_history": ["Sift"],
                    "education": ["Mathematics"],
                    "experience_years": 4,
                }
            )
        if "Python experience" in prompt:
            return json.dumps(
                {"query_id": "Q1", "status": "insufficient_evidence", "confidence": 0.2, "evidence": ""}
            )
        evidence = "shipped a policy-compliant hiring platform"
        query_id = "Q2"
        for index in range(1, 11):
            marker = f"Query: Q{index} "
            if marker in prompt:
                query_id = f"Q{index}"
                break
        return json.dumps({"query_id": query_id, "status": "met", "confidence": 0.9, "evidence": evidence})


class LocalSkillProvider:
    def __init__(self, *, fail: bool = False) -> None:
        self.calls = 0
        self.prompts: list[str] = []
        self.fail = fail

    async def complete(self, prompt: str, *, system: str | None = None) -> str:
        del system
        self.calls += 1
        self.prompts.append(prompt)
        if self.fail:
            raise RuntimeError("ollama down")
        return json.dumps(
            {
                "reviews": [
                    {
                        "query_id": "Q1",
                        "status": "partially_met",
                        "evidence": "expertise in FastAPI",
                    }
                ]
            }
        )


def _python_checklist() -> list[dict[str, object]]:
    checklist = sample_queries()
    checklist[0] = {
        **checklist[0],
        "question": "Does the candidate have Python experience?",
        "dimension": "technical",
        "evidence_hints": ["Python"],
    }
    return checklist


@pytest.mark.asyncio
async def test_local_model_raises_implied_python_without_touching_met_items() -> None:
    cloud = CloudScreenProvider()
    local = LocalSkillProvider()
    run = build_cv_screening_graph(FallbackLLMManager(cloud, cloud, local), HashingEmbedder())
    result = await run(
        {
            "filename": "ada.docx",
            "file_bytes": _docx_with(PYTHON_CV),
            "job_description": "A" * 120,
            "locked_checklist": _python_checklist(),
        }
    )
    by_query = {item["query_id"]: item for item in result["query_evaluations"]}
    assert by_query["Q1"]["status"] == "partially_met"
    assert "FastAPI" in str(by_query["Q1"]["evidence"])
    assert by_query["Q2"]["status"] == "met"
    assert cloud.calls == 11
    assert local.calls == 1
    assert "Python experience" in local.prompts[0]
    assert "criterion 2" not in local.prompts[0]


@pytest.mark.asyncio
async def test_local_model_failure_keeps_the_original_scores() -> None:
    cloud = CloudScreenProvider()
    local = LocalSkillProvider(fail=True)
    run = build_cv_screening_graph(FallbackLLMManager(cloud, cloud, local), HashingEmbedder())
    result = await run(
        {
            "filename": "ada.docx",
            "file_bytes": _docx_with(PYTHON_CV),
            "job_description": "A" * 120,
            "locked_checklist": _python_checklist(),
        }
    )
    by_query = {item["query_id"]: item for item in result["query_evaluations"]}
    assert by_query["Q1"]["status"] == "insufficient_evidence"
    assert by_query["Q2"]["status"] == "met"
    assert local.calls == 1
