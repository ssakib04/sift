from io import BytesIO

import pytest
from docx import Document

from app.services.documents import EmptyDocumentError, UnsupportedDocumentError, chunk_text, extract_text


def test_chunk_text_overlaps_long_input() -> None:
    text = "a" * 1000
    chunks = chunk_text(text, size=800, overlap=100)
    assert chunks[0] == "a" * 800
    assert chunks[1].startswith("a")
    assert len(chunks) == 2
    assert "".join(chunks)[100:800] == text[100:800]


def test_extract_docx_text() -> None:
    document = Document()
    document.add_paragraph("Remote work requires manager approval.")
    buffer = BytesIO()
    document.save(buffer)
    assert extract_text("policy.docx", buffer.getvalue()) == "Remote work requires manager approval."


def test_extract_rejects_unknown_types() -> None:
    with pytest.raises(UnsupportedDocumentError):
        extract_text("notes.txt", b"hello")


def test_extract_rejects_empty_docx() -> None:
    document = Document()
    buffer = BytesIO()
    document.save(buffer)
    with pytest.raises(EmptyDocumentError):
        extract_text("empty.docx", buffer.getvalue())
